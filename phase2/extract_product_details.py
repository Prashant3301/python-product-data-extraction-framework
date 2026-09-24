import os
import re
import time
import random
from datetime import datetime

import pandas as pd
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.edge.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.common.exceptions import InvalidSessionIdException, WebDriverException, TimeoutException
from webdriver_manager.microsoft import EdgeChromiumDriverManager
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SCRIPT_FOLDER = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(SCRIPT_FOLDER, "all_medicine_urls.xlsx")
OUTPUT_FILE = os.path.join(SCRIPT_FOLDER, "1mg_medicine_complete_details.xlsx")
CHECKPOINT_FILE = os.path.join(SCRIPT_FOLDER, "1mg_medicine_details_checkpoint.xlsx")

MAX_MEDICINES = None
SAVE_AFTER_EVERY = 5
RESTART_DRIVER_AFTER = 25
MIN_DELAY_SECONDS = 5
MAX_DELAY_SECONDS = 8
MAX_RETRIES = 5
HEADLESS_MODE = False


def clean_text(value):
    if value is None:
        return ""
    value = str(value).replace("\xa0", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n\s*\n+", "\n", value)
    return value.strip()


def single_line(value):
    value = clean_text(value)
    value = re.sub(r"\s*\n\s*", " | ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" |")


def first_match(pattern, text, flags=re.IGNORECASE | re.DOTALL):
    match = re.search(pattern, text, flags)
    return clean_text(match.group(1)) if match else ""


def get_text_between(text, start_pattern, end_patterns):
    end_pattern = "|".join(end_patterns)
    pattern = start_pattern + r"\s*(.*?)" + r"(?=\n(?:" + end_pattern + r")|\Z)"
    return single_line(first_match(pattern, text))


def load_medicine_urls():
    if not os.path.exists(INPUT_FILE):
        raise FileNotFoundError("Input Excel file was not found:\n" + INPUT_FILE)

    dataframe = pd.read_excel(INPUT_FILE, engine="openpyxl")
    print("\nInput Excel columns:")
    print(list(dataframe.columns))

    link_column = None
    for column in dataframe.columns:
        if str(column).strip().lower() in ("link", "url", "medicine url"):
            link_column = column
            break

    if link_column is None:
        raise ValueError("The input Excel must contain a Link, URL, or Medicine URL column.")

    raw_urls = dataframe[link_column].dropna().astype(str).str.strip().tolist()
    medicine_urls = []
    for url in raw_urls:
        if url.lower().startswith("https://www.1mg.com/drugs/"):
            medicine_urls.append(url.split("?")[0].split("#")[0])

    medicine_urls = list(dict.fromkeys(medicine_urls))
    return medicine_urls if MAX_MEDICINES is None else medicine_urls[:MAX_MEDICINES]


def create_edge_driver():
    options = webdriver.EdgeOptions()
    if HEADLESS_MODE:
        options.add_argument("--headless=new")
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-popup-blocking")
    options.add_argument("--disable-extensions")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--no-sandbox")
    options.add_argument("--log-level=3")

    service = Service(EdgeChromiumDriverManager().install())
    driver = webdriver.Edge(service=service, options=options)
    driver.set_page_load_timeout(60)
    return driver


def close_driver_safely(driver):
    if driver is not None:
        try:
            driver.quit()
        except Exception:
            pass


def restart_edge_driver(driver):
    print("Edge session unavailable. Restarting Edge...")
    close_driver_safely(driver)
    time.sleep(3)
    new_driver = create_edge_driver()
    print("New Edge session started.")
    return new_driver


def driver_is_alive(driver):
    if driver is None:
        return False
    try:
        driver.execute_script("return 1;")
        return True
    except Exception:
        return False


def extract_medicine_name(soup):
    heading = soup.find("h1")
    if heading:
        name = clean_text(heading.get_text(" ", strip=True))
        if name:
            return name
    if soup.title:
        title = clean_text(soup.title.get_text(" ", strip=True))
        return re.sub(r"\s*[-|]\s*(Tata\s*)?1mg.*$", "", title, flags=re.IGNORECASE)
    return ""


def extract_composition(page_text):
    return single_line(first_match(r"Composition:\s*(.*?)\s*Marketer details:", page_text))


def extract_manufacturer(page_text):
    value = first_match(
        r"Marketer details:\s*(.*?)(?=\n(?:Order Medicines|Genuine|Authenticity Assured|"
        r"NPPA Regulated|Payment, Returns|Price Info|Product introduction|\d+%\s+cheaper alternative|\Z))",
        page_text,
    )
    value = single_line(value)
    value = re.sub(r"\s*\|\s*Product Images.*$", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\s*\|\s*\+\d+\s*more", "", value, flags=re.IGNORECASE)
    return value.strip(" |")[:300]


def extract_pack_size(page_text, medicine_name):
    search_text = page_text
    if medicine_name:
        position = page_text.lower().find(medicine_name.lower())
        if position >= 0:
            search_text = page_text[position:position + 1500]

    patterns = (
        r"\b(strip of \d+\s+[^\n]+)", r"\b(bottle of \d+(?:\.\d+)?\s*[^\n]+)",
        r"\b(vial of \d+\s+[^\n]+)", r"\b(prefilled syringe of [^\n]+)",
        r"\b(tube of \d+(?:\.\d+)?\s*[^\n]+)", r"\b(packet of \d+(?:\.\d+)?\s*[^\n]+)",
        r"\b(box of \d+(?:\.\d+)?\s*[^\n]+)", r"\b(\d+\s+tablets)\b",
        r"\b(\d+\s+capsules)\b", r"\b(\d+\s+injections?)\b",
    )
    for pattern in patterns:
        value = first_match(pattern, search_text)
        if value:
            return single_line(value)
    return ""


def extract_prescription_status(page_text):
    patterns = (r"\bPrescription Required\b", r"\bPrescription is required\b", r"\bRequires prescription\b")
    return "Yes" if any(re.search(p, page_text, re.IGNORECASE) for p in patterns) else "Not shown"


def extract_prices(page_text):
    mrp = first_match(r"\bMRP\s*₹\s*([\d,]+(?:\.\d+)?)", page_text)
    discount = first_match(r"\b(\d+(?:\.\d+)?%\s*off)\b", page_text)
    selling_price = ""

    mrp_match = re.search(r"\bMRP\s*₹\s*[\d,]+(?:\.\d+)?", page_text, flags=re.IGNORECASE)
    if mrp_match:
        preceding_text = page_text[max(0, mrp_match.start() - 300):mrp_match.start()]
        preceding_prices = re.findall(r"₹\s*([\d,]+(?:\.\d+)?)", preceding_text)
        if preceding_prices:
            selling_price = preceding_prices[-1]

    if not selling_price:
        all_prices = re.findall(r"₹\s*([\d,]+(?:\.\d+)?)", page_text)
        if all_prices:
            selling_price = all_prices[0]

    offer_prices = re.findall(r"Get for\s*₹\s*([\d,]+(?:\.\d+)?)", page_text, flags=re.IGNORECASE)
    offer_prices = list(dict.fromkeys(offer_prices))

    return {
        "Selling Price (INR)": selling_price,
        "MRP (INR)": mrp,
        "Discount": discount,
        "Offer Prices (INR)": ", ".join(offer_prices),
    }


def extract_availability(page_text):
    if re.search(r"\bNOT AVAILABLE\b", page_text, re.IGNORECASE):
        return "Not available"
    if re.search(r"\bADD\b", page_text, re.IGNORECASE):
        return "Available"
    return "Not confirmed"


def section(page_text, heading, ends):
    return get_text_between(page_text, heading, ends)


def extract_product_introduction(page_text):
    value = section(page_text, r"Product introduction", (
        r"Uses of [^\n]+", r"Benefits of [^\n]+", r"Side effects of [^\n]+",
        r"How to use [^\n]+", r"How [^\n]+ works", r"Quick tips", r"Safety advice", r"Fact box",
    ))
    return single_line(re.sub(r"^Product Summary\s*:\s*", "", value, flags=re.IGNORECASE))


def extract_uses(t):
    return section(t, r"Uses of [^\n]+", (r"Benefits of [^\n]+", r"Side effects of [^\n]+", r"How to use [^\n]+", r"How [^\n]+ works", r"Quick tips", r"Safety advice"))


def extract_benefits(t):
    return section(t, r"Benefits of [^\n]+", (r"Side effects of [^\n]+", r"How to use [^\n]+", r"How [^\n]+ works", r"Quick tips", r"Safety advice"))


def extract_side_effects(t):
    return section(t, r"Side effects of [^\n]+", (r"How to use [^\n]+", r"How [^\n]+ works", r"Quick tips", r"Safety advice", r"Fact box"))


def extract_how_to_use(t):
    return section(t, r"How to use [^\n]+", (r"How [^\n]+ works", r"Quick tips", r"Safety advice", r"Fact box"))


def extract_how_it_works(t):
    return section(t, r"How [^\n]+ works", (r"Quick tips", r"Safety advice", r"Fact box", r"Interaction with drugs"))


def extract_quick_tips(t):
    return section(t, r"Quick tips", (r"Safety advice", r"Fact box", r"Interaction with drugs", r"Pregnancy", r"Breast feeding", r"Driving", r"Kidney", r"Liver", r"₹\s*[\d,.]+", r"Save more with additional offers"))


def parse_medicine_page(html, url):
    soup = BeautifulSoup(html, "html.parser")
    page_text = clean_text(soup.get_text(separator="\n", strip=True))
    name = extract_medicine_name(soup)
    prices = extract_prices(page_text)
    return {
        "Medicine Name": name,
        "Prescription Required": extract_prescription_status(page_text),
        "Pack Size": extract_pack_size(page_text, name),
        "Composition": extract_composition(page_text),
        "Manufacturer / Marketer": extract_manufacturer(page_text),
        "Selling Price (INR)": prices["Selling Price (INR)"],
        "MRP (INR)": prices["MRP (INR)"],
        "Discount": prices["Discount"],
        "Offer Prices (INR)": prices["Offer Prices (INR)"],
        "Availability": extract_availability(page_text),
        "Product Introduction": extract_product_introduction(page_text),
        "Uses": extract_uses(page_text),
        "Benefits": extract_benefits(page_text),
        "Side Effects": extract_side_effects(page_text),
        "How to Use": extract_how_to_use(page_text),
        "How It Works": extract_how_it_works(page_text),
        "Quick Tips": extract_quick_tips(page_text),
        "Expiry Information": first_match(r"(Product expires after\s+[^\n]+)", page_text),
        "Bought Recently": first_match(r"([\d,]+\s+people)\s+bought recently", page_text),
        "Medicine URL": url,
        "Scrape Status": "Success",
        "Error Message": "",
        "Scraped At": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def create_failed_record(url, error_message):
    columns = ["Medicine Name", "Prescription Required", "Pack Size", "Composition", "Manufacturer / Marketer",
               "Selling Price (INR)", "MRP (INR)", "Discount", "Offer Prices (INR)", "Availability",
               "Product Introduction", "Uses", "Benefits", "Side Effects", "How to Use", "How It Works",
               "Quick Tips", "Expiry Information", "Bought Recently"]
    row = {column: "" for column in columns}
    row.update({"Medicine URL": url, "Scrape Status": "Failed", "Error Message": error_message,
                "Scraped At": datetime.now().strftime("%Y-%m-%d %H:%M:%S")})
    return row


def save_excel(records, filename):
    if not records:
        return
    dataframe = pd.DataFrame(records)
    if "Medicine URL" in dataframe.columns:
        dataframe = dataframe.drop_duplicates(subset=["Medicine URL"], keep="last")
    dataframe.to_excel(filename, index=False, engine="openpyxl")

    workbook = load_workbook(filename)
    worksheet = workbook.active
    worksheet.title = "Medicine Details"
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    widths = {"Medicine Name": 35, "Composition": 55, "Manufacturer / Marketer": 40,
              "Product Introduction": 70, "Uses": 60, "Benefits": 70, "Side Effects": 60,
              "How to Use": 60, "How It Works": 60, "Quick Tips": 70, "Medicine URL": 70,
              "Error Message": 45}
    for number, cell in enumerate(worksheet[1], start=1):
        worksheet.column_dimensions[get_column_letter(number)].width = widths.get(cell.value, 22)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(filename)


def load_checkpoint():
    if not os.path.exists(CHECKPOINT_FILE):
        return [], set()
    try:
        df = pd.read_excel(CHECKPOINT_FILE, engine="openpyxl").fillna("")
        if "Scrape Status" in df.columns:
            df = df[df["Scrape Status"].astype(str).str.strip().str.lower().eq("success")].copy()
        results = df.to_dict(orient="records")
        processed = set(df["Medicine URL"].astype(str).str.strip().tolist()) if "Medicine URL" in df.columns else set()
        print(f"\nCheckpoint loaded. Successful records found: {len(processed)}")
        return results, processed
    except Exception as error:
        print("\nCheckpoint could not be loaded:", error)
        return [], set()


def scrape_url_with_retries(driver, url):
    final_error = ""
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            if not driver_is_alive(driver):
                driver = restart_edge_driver(driver)
            driver.get(url)
            WebDriverWait(driver, 30).until(lambda d: d.execute_script("return document.readyState") == "complete")
            time.sleep(3)
            for fraction in (0.40, 0.75, 1.0):
                driver.execute_script(f"window.scrollTo(0, document.body.scrollHeight * {fraction});")
                time.sleep(1)
            html = driver.page_source
            if len(html) < 1000:
                raise ValueError("Returned HTML was unexpectedly short.")
            record = parse_medicine_page(html, url)
            if not record["Medicine Name"]:
                raise ValueError("Medicine name could not be extracted.")
            return driver, record, ""
        except (InvalidSessionIdException, TimeoutException) as error:
            final_error = str(error)
            print(f"Attempt {attempt}: browser session ended or timed out.")
            driver = restart_edge_driver(driver)
        except WebDriverException as error:
            final_error = str(error)
            print(f"Attempt {attempt} WebDriver error: {final_error}")
            driver = restart_edge_driver(driver)
        except Exception as error:
            final_error = str(error)
            print(f"Attempt {attempt} failed: {final_error}")
        if attempt < MAX_RETRIES:
            time.sleep(3)
    return driver, None, final_error


def main():
    print("=" * 70)
    print("PHASE 2: ROBUST 1MG MEDICINE DETAILS SCRAPER")
    print("=" * 70)
    urls = load_medicine_urls()
    print(f"\nValid medicine URLs found: {len(urls)}")
    if not urls:
        print("No valid medicine URLs were found.")
        return

    results, processed_urls = load_checkpoint()
    driver = None
    completed_this_run = 0
    try:
        driver = create_edge_driver()
        for index, url in enumerate(urls, start=1):
            if url in processed_urls:
                print(f"[{index}/{len(urls)}] Already processed successfully. Skipping.")
                continue
            print(f"\n[{index}/{len(urls)}] Opening:\n{url}")
            driver, record, error = scrape_url_with_retries(driver, url)
            if record is not None:
                print("Saved:", record["Medicine Name"])
                results.append(record)
                processed_urls.add(url)
                completed_this_run += 1
            else:
                print("Failed after all retries:", url)
                results.append(create_failed_record(url, error))

            if completed_this_run > 0 and completed_this_run % SAVE_AFTER_EVERY == 0:
                save_excel(results, CHECKPOINT_FILE)
                print(f"Checkpoint saved. Successful URLs: {len(processed_urls)}")
            if completed_this_run > 0 and completed_this_run % RESTART_DRIVER_AFTER == 0:
                save_excel(results, CHECKPOINT_FILE)
                driver = restart_edge_driver(driver)
            time.sleep(random.uniform(MIN_DELAY_SECONDS, MAX_DELAY_SECONDS))
    except KeyboardInterrupt:
        print("\nStopped manually. Saving progress...")
        save_excel(results, CHECKPOINT_FILE)
    finally:
        close_driver_safely(driver)

    save_excel(results, OUTPUT_FILE)
    successes = sum(r.get("Scrape Status") == "Success" for r in results)
    failures = sum(r.get("Scrape Status") == "Failed" for r in results)
    print("\n" + "=" * 70)
    print("PHASE 2 FINISHED")
    print("=" * 70)
    print(f"Successful records: {successes}")
    print(f"Failed records: {failures}")
    print("Final Excel file:", OUTPUT_FILE)
    if failures == 0 and os.path.exists(CHECKPOINT_FILE):
        try:
            os.remove(CHECKPOINT_FILE)
        except OSError:
            pass


if __name__ == "__main__":
    main()
