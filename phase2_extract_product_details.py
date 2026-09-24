"""
Phase 2: Product Detail Extraction Demo

This module demonstrates a resilient product-data extraction pipeline
using locally stored HTML fixture files.

The project intentionally does not target a live third-party website.
Only use automation against systems where you have authorization and
where automated access is permitted by the applicable terms.

Demonstrated engineering concepts:
- HTML parsing with BeautifulSoup
- Configurable extraction rules
- Retry handling
- Checkpoint persistence and resume
- Duplicate removal
- Structured Excel export
- Error tracking
"""

import re
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
INPUT_FILE = PROJECT_ROOT / "product_sources.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "product_details.xlsx"
CHECKPOINT_FILE = PROJECT_ROOT / "product_details_checkpoint.xlsx"

SAVE_AFTER_EVERY = 5
MAX_RETRIES = 3
RETRY_DELAY_SECONDS = 2


# ============================================================
# TEXT HELPERS
# ============================================================

def clean_text(value):
    """Normalize spaces and repeated line breaks."""

    if value is None:
        return ""

    value = str(value).replace("\xa0", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n\s*\n+", "\n", value)

    return value.strip()


def single_line(value):
    """Convert multiline text into one Excel-friendly line."""

    value = clean_text(value)
    value = re.sub(r"\s*\n\s*", " | ", value)
    value = re.sub(r"\s+", " ", value)

    return value.strip(" |")


def first_match(pattern, text, flags=re.IGNORECASE | re.DOTALL):
    """Return the first regular-expression capture group."""

    match = re.search(pattern, text, flags)

    if match:
        return clean_text(match.group(1))

    return ""


def get_text_between(text, start_pattern, end_patterns):
    """Extract text between a starting heading and the next heading."""

    if end_patterns:
        end_pattern = "|".join(end_patterns)
        pattern = (
            start_pattern
            + r"\s*(.*?)"
            + r"(?=\n(?:"
            + end_pattern
            + r")|\Z)"
        )
    else:
        pattern = start_pattern + r"\s*(.*)\Z"

    return single_line(first_match(pattern, text))


# ============================================================
# INPUT
# ============================================================

def load_product_sources():
    """
    Load local HTML fixture paths from product_sources.xlsx.

    Required columns:
    - Product Name
    - HTML File
    """

    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Input file was not found: {INPUT_FILE}"
        )

    dataframe = pd.read_excel(
        INPUT_FILE,
        engine="openpyxl",
    )

    required_columns = {
        "Product Name",
        "HTML File",
    }

    missing_columns = required_columns - set(dataframe.columns)

    if missing_columns:
        raise ValueError(
            "Missing input columns: "
            + ", ".join(sorted(missing_columns))
        )

    records = []

    for _, row in dataframe.iterrows():
        html_file = clean_text(row.get("HTML File"))

        if not html_file:
            continue

        file_path = Path(html_file)

        if not file_path.is_absolute():
            file_path = PROJECT_ROOT / file_path

        records.append({
            "Input Product Name": clean_text(
                row.get("Product Name")
            ),
            "HTML File": str(file_path.resolve()),
        })

    return records


# ============================================================
# EXTRACTION
# ============================================================

def extract_product_name(soup):
    """Extract a product name from H1 or the HTML title."""

    heading = soup.find("h1")

    if heading:
        value = clean_text(heading.get_text(" ", strip=True))

        if value:
            return value

    if soup.title:
        return clean_text(soup.title.get_text(" ", strip=True))

    return ""


def extract_composition(page_text):
    """Extract composition from the sample page."""

    return single_line(
        first_match(
            r"Composition:\s*(.*?)\s*Manufacturer:",
            page_text,
        )
    )


def extract_manufacturer(page_text):
    """Extract manufacturer from the sample page."""

    return single_line(
        first_match(
            r"Manufacturer:\s*(.*?)"
            r"(?=\n(?:Pack Size:|Price:|Product Introduction:|\Z))",
            page_text,
        )
    )


def extract_pack_size(page_text):
    """Extract pack size from the sample page."""

    return single_line(
        first_match(
            r"Pack Size:\s*([^\n]+)",
            page_text,
        )
    )


def extract_price(page_text):
    """Extract a sample price displayed in INR."""

    return first_match(
        r"Price:\s*₹\s*([\d,]+(?:\.\d+)?)",
        page_text,
    )


def extract_product_introduction(page_text):
    """Extract the product-introduction section."""

    return get_text_between(
        page_text,
        r"Product Introduction:",
        (
            r"Uses:",
            r"Benefits:",
            r"Side Effects:",
            r"How to Use:",
            r"How It Works:",
            r"Quick Tips:",
        ),
    )


def extract_section(page_text, heading, end_headings):
    """Extract a named section from a sample product page."""

    return get_text_between(
        page_text,
        re.escape(heading),
        tuple(re.escape(item) for item in end_headings),
    )


def parse_product_file(source):
    """Read and parse one locally stored HTML fixture."""

    html_path = Path(source["HTML File"])

    if not html_path.exists():
        raise FileNotFoundError(
            f"HTML fixture was not found: {html_path}"
        )

    html = html_path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")

    page_text = clean_text(
        soup.get_text(separator="\n", strip=True)
    )

    product_name = extract_product_name(soup)

    if not product_name:
        product_name = source["Input Product Name"]

    return {
        "Product Name": product_name,
        "Composition": extract_composition(page_text),
        "Manufacturer": extract_manufacturer(page_text),
        "Pack Size": extract_pack_size(page_text),
        "Price (INR)": extract_price(page_text),
        "Product Introduction": extract_product_introduction(
            page_text
        ),
        "Uses": extract_section(
            page_text,
            "Uses:",
            (
                "Benefits:",
                "Side Effects:",
                "How to Use:",
                "How It Works:",
                "Quick Tips:",
            ),
        ),
        "Benefits": extract_section(
            page_text,
            "Benefits:",
            (
                "Side Effects:",
                "How to Use:",
                "How It Works:",
                "Quick Tips:",
            ),
        ),
        "Side Effects": extract_section(
            page_text,
            "Side Effects:",
            (
                "How to Use:",
                "How It Works:",
                "Quick Tips:",
            ),
        ),
        "How to Use": extract_section(
            page_text,
            "How to Use:",
            (
                "How It Works:",
                "Quick Tips:",
            ),
        ),
        "How It Works": extract_section(
            page_text,
            "How It Works:",
            ("Quick Tips:",),
        ),
        "Quick Tips": extract_section(
            page_text,
            "Quick Tips:",
            (),
        ),
        "Source File": str(html_path),
        "Extraction Status": "Success",
        "Error Message": "",
        "Extracted At": datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
    }


def create_failed_record(source, error_message):
    """Create a structured record for a failed extraction."""

    return {
        "Product Name": source.get("Input Product Name", ""),
        "Composition": "",
        "Manufacturer": "",
        "Pack Size": "",
        "Price (INR)": "",
        "Product Introduction": "",
        "Uses": "",
        "Benefits": "",
        "Side Effects": "",
        "How to Use": "",
        "How It Works": "",
        "Quick Tips": "",
        "Source File": source.get("HTML File", ""),
        "Extraction Status": "Failed",
        "Error Message": error_message,
        "Extracted At": datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
    }


# ============================================================
# CHECKPOINT
# ============================================================

def load_checkpoint():
    """Load successfully processed local fixture files."""

    if not CHECKPOINT_FILE.exists():
        return [], set()

    try:
        dataframe = pd.read_excel(
            CHECKPOINT_FILE,
            engine="openpyxl",
        ).fillna("")

        if "Extraction Status" in dataframe.columns:
            dataframe = dataframe[
                dataframe["Extraction Status"]
                .astype(str)
                .str.strip()
                .str.lower()
                .eq("success")
            ].copy()

        records = dataframe.to_dict(orient="records")
        processed_files = set()

        if "Source File" in dataframe.columns:
            processed_files = set(
                dataframe["Source File"]
                .astype(str)
                .str.strip()
                .tolist()
            )

        print(
            f"Checkpoint loaded: "
            f"{len(processed_files)} successful records"
        )

        return records, processed_files

    except Exception as error:
        print(f"Checkpoint could not be loaded: {error}")
        return [], set()


# ============================================================
# EXCEL OUTPUT
# ============================================================

def save_excel(records, filename):
    """Save results to a formatted Excel workbook."""

    if not records:
        return

    dataframe = pd.DataFrame(records)

    if "Source File" in dataframe.columns:
        dataframe = dataframe.drop_duplicates(
            subset=["Source File"],
            keep="last",
        )

    dataframe.to_excel(
        filename,
        index=False,
        engine="openpyxl",
    )

    workbook = load_workbook(filename)
    worksheet = workbook.active
    worksheet.title = "Product Details"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 35

    column_widths = {
        "Product Name": 35,
        "Composition": 50,
        "Manufacturer": 35,
        "Pack Size": 25,
        "Price (INR)": 15,
        "Product Introduction": 70,
        "Uses": 55,
        "Benefits": 55,
        "Side Effects": 55,
        "How to Use": 55,
        "How It Works": 55,
        "Quick Tips": 55,
        "Source File": 65,
        "Extraction Status": 18,
        "Error Message": 45,
        "Extracted At": 22,
    }

    for column_number, cell in enumerate(
        worksheet[1],
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = column_widths.get(cell.value, 22)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    workbook.save(filename)


# ============================================================
# RETRY HANDLING
# ============================================================

def extract_with_retries(source):
    """Retry local file extraction after temporary failures."""

    final_error = ""

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            record = parse_product_file(source)
            return record, ""

        except Exception as error:
            final_error = str(error)

            print(
                f"Attempt {attempt}/{MAX_RETRIES} failed: "
                f"{final_error}"
            )

            if attempt < MAX_RETRIES:
                time.sleep(RETRY_DELAY_SECONDS)

    return None, final_error


# ============================================================
# MAIN
# ============================================================

def main():
    """Run the local demonstration extraction pipeline."""

    print("=" * 70)
    print("PHASE 2: PRODUCT DETAIL EXTRACTION DEMO")
    print("=" * 70)

    sources = load_product_sources()

    print(
        f"Input fixture files found: {len(sources)}"
    )

    results, processed_files = load_checkpoint()
    completed_this_run = 0

    try:
        for index, source in enumerate(sources, start=1):
            source_file = source["HTML File"]

            if source_file in processed_files:
                print(
                    f"[{index}/{len(sources)}] "
                    "Already processed. Skipping."
                )
                continue

            print(
                f"\n[{index}/{len(sources)}] "
                f"Processing: {source_file}"
            )

            record, error_message = extract_with_retries(source)

            if record is not None:
                results.append(record)
                processed_files.add(source_file)
                completed_this_run += 1

                print(f"Saved: {record['Product Name']}")

            else:
                results.append(
                    create_failed_record(
                        source,
                        error_message,
                    )
                )
                print("Failed after all retry attempts.")

            if (
                completed_this_run > 0
                and completed_this_run % SAVE_AFTER_EVERY == 0
            ):
                save_excel(results, CHECKPOINT_FILE)
                print("Checkpoint saved.")

    except KeyboardInterrupt:
        print("\nStopped manually. Saving progress...")
        save_excel(results, CHECKPOINT_FILE)

    save_excel(results, OUTPUT_FILE)

    successes = sum(
        record.get("Extraction Status") == "Success"
        for record in results
    )
    failures = sum(
        record.get("Extraction Status") == "Failed"
        for record in results
    )

    print("\n" + "=" * 70)
    print("EXTRACTION FINISHED")
    print("=" * 70)
    print(f"Successful records: {successes}")
    print(f"Failed records: {failures}")
    print(f"Output: {OUTPUT_FILE}")

    if failures == 0 and CHECKPOINT_FILE.exists():
        try:
            CHECKPOINT_FILE.unlink()
            print("Checkpoint removed after successful completion.")
        except OSError:
            pass


if __name__ == "__main__":
    main()
